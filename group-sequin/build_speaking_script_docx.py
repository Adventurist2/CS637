"""Create the three-speaker speaking script for combined slides 15 to 30."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output'/'documents'/'SEQUIN_15_Minute_Three_Speaker_Script.docx'
PARTS=[
('Part One Critique and Line Protection','Speaker 1  |  Slides 15 to 19  |  0:00 to 5:00',[
('Slide 15 Limitations and Three Extensions',[
"Now that we have explained SEQUIN and its implementation, we move to the second part of our project.",
"The paper helps us identify dangerous sequences of transmission-line failures. We wanted to investigate how this analysis could support decisions about making the grid more resilient.",
"We identified three practical gaps. First, finding a dangerous attack does not directly tell us which lines we should protect when our budget is limited. Second, failures may arrive at different speeds. The operator may have more time to respond between some failures and very little time between others. Third, maximizing total electricity supplied treats every megawatt equally. In practice, some services may deserve higher priority.",
"These gaps motivated our three extensions: line protection, variable response windows, and critical-service priorities."]),
('Slide 16 Scenario Aware Line Protection',[
"Our first extension asks a simple question: if we can protect only a few lines, which lines should we choose?",
"The important point is that the attack can change after protection. Suppose we protect a line from the original worst attack. Another sequence might then become the most damaging attack. Therefore, replaying only the original attack would give an incomplete assessment.",
"We evaluate every eligible ordered pair of failures on our selected small networks. Our method builds the protection plan one line at a time. At each step, it chooses the additional protected line that gives the smallest worst loss across the design scenarios.",
"Here, protection means the line cannot fail through the particular failure mechanism we model. We assume equal protection costs."]),
('Slide 17 Protection Evaluation Protocol',[
"We evaluated this extension on cases 14 and 39. For two sequential failures, case 14 has 380 possible ordered pairs, while case 39 has 2,070.",
"We separated the demand scenarios used to select protection from those used to evaluate it. This allows us to check whether a plan also helps under unseen demand conditions.",
"We compared our method with several alternatives, including random protection, protecting high-flow lines, and protection designed using only one demand scenario.",
"We also checked small protection budgets through complete enumeration. This helps us distinguish the quality of our protection-selection method from the quality of the attack search."]),
('Slide 18 Protection Results',[
"On these plots, lower load loss is better.",
"For case 14, protecting two lines using our scenario-based method reduces mean held-out worst loss from about 95 percent without protection to about 85 percent. At the same budget, initial-flow protection leaves about 93 percent loss. So, on this case, our selected protection plan provides a meaningful improvement.",
"However, the result is mixed on case 39. At a budget of two lines, our method leaves about 20.47 percent loss, while initial-flow protection leaves about 20.35 percent. The simpler baseline performs slightly better.",
"This means our method can find useful plans, but we cannot claim that it always chooses the best plan for unseen conditions."]),
('Slide 19 Robustness',[
"We then changed demand conditions and generator response limits while keeping the protection plans fixed. Our nominal robust plan performed worse than flow protection in 10 of 36 feasible stress comparisons.",
"Some stress scenarios were already infeasible before an attack, so we recorded them separately.",
"The conclusion is that protection helps, but its success depends on the conditions used during design. A useful next step is to design protection across both demand uncertainty and response-speed uncertainty.",
"I will now hand over to the next speaker to explain the timing and critical-service extensions."])
]),
('Part Two Timing and Critical Services','Speaker 2  |  Slides 20 to 24  |  5:00 to 10:00',[
('Slide 20 Variable Response Windows',[
"Our second extension studies the time available between failures.",
"The basic idea is that generators cannot change their output instantly. A longer response window allows more adjustment, while a shorter window allows less.",
"We compare three patterns over the same total horizon of 15 minutes. The even pattern uses three five-minute windows. The early-cluster pattern uses one minute, one minute, and thirteen minutes. The late-cluster pattern uses thirteen minutes, one minute, and one minute.",
"We also make the sequential and simultaneous comparisons fairer by giving both three redispatch opportunities. For the simultaneous case, all selected lines fail in the first window, but the operator can still adjust again during the remaining windows."]),
('Slide 21 Timing Results',[
"This slide shows that final load loss and time-weighted loss answer different questions. The left plot measures how much demand is unserved at the end.",
"The right plot measures a snapshot energy proxy. We calculate it by multiplying the loss at each response snapshot by the duration of its window, then adding the contributions.",
"For case 118, the sequential even pattern gives approximately 26 percent final loss. The late-cluster pattern gives approximately 35 percent final loss, but its energy proxy is much smaller. This happens because the larger loss appears in the short later windows and receives less time weight.",
"Therefore, a smaller energy proxy does not necessarily mean a better final condition. We also clearly label this measure as a proxy. We do not simulate the continuous generator and service trajectories inside each window."]),
('Slide 22 Critical Service Objectives',[
"Our third extension changes which demand the operator tries hardest to preserve. The original objective maximizes total served electricity. It does not distinguish between different service types.",
"We introduce synthetic priority groups to represent the idea of critical services. These are illustrative groups, not identified hospitals or other real facilities.",
"We compare four main objectives. The total-service baseline values every megawatt equally. Weight two gives critical demand twice the importance of ordinary demand. Weight ten gives it ten times the importance. Critical-first first maximizes critical service, then maximizes total service while retaining that critical-service optimum within a small numerical tolerance.",
"Every policy starts from the same initial operating condition and faces the same evaluated attack sequences."]),
('Slide 23 Critical Service Results',[
"Here, higher critical-service retention is better.",
"At the nominal setting, the total-service baseline retains about 67 percent of critical demand in case 39 and about 77 percent in case 118, using our worst-case summary over the evaluated attacks. The priority policies increase the measured retention to approximately 100 percent in these larger cases.",
"However, this is a result for the evaluated attack cohort. It is not a guarantee against every possible attack.",
"Case 14 shows a much smaller improvement. This reminds us that changing priorities cannot remove the underlying network constraints."]),
('Slide 24 Gains and Adverse Cases',[
"Priority service also has costs. In one deliberately selected positive example, critical-first preserves about 1,573 megawatts more critical demand than the baseline, while increasing total loss by about 121 megawatts. This means some ordinary demand is sacrificed to preserve critical demand.",
"We also found an adverse example where critical-first causes about 240 megawatts more critical loss after the later failure.",
"The operator optimizes the current response without knowing the next outage. An early decision can therefore create a worse later state.",
"These examples are selected extremes, not average outcomes. Our conclusion is that priorities are useful, but they require evaluation over complete failure sequences.",
"The final speaker will explain our checks, demonstrations and conclusions."])
]),
('Part Three Verification Demonstration and Conclusions','Speaker 3  |  Slides 25 to 30  |  10:00 to 15:00  |  Easiest role',[
('Slide 25 Sensitivity to Response Speed',[
"We also checked whether the conclusions change when generators have different response limits.",
"The horizontal axis represents the ramp allowance: how much a generator can change its output during a response. The vertical axis shows the critical demand retained under the evaluated attacks.",
"The priority policies generally preserve more critical service than the total-service baseline in these selected larger-grid experiments. However, these curves describe a shared sampled attack cohort. They do not prove that all possible attacks have been covered.",
"We also tested different sizes of synthetic critical groups. Changing the group changes both how many buses receive priority and which buses receive it."]),
('Slide 26 Correctness and Limitations',[
"We used several checks to make the results reviewable. First, all 16 extension tests passed.",
"Second, we compared the extended model with the untouched group implementation on 26 state probes. The generation and served-load values matched exactly.",
"Third, we independently reran complete post-protection attack searches on selected held-out cases. Their worst losses matched the saved outage-table results. We also checked that each policy faced the same attack sequences within its comparison.",
"There is a separate limitation in the original reproduction benchmark. Its checker fails because a case-162 result at eight failures is missing after an infeasible response. We preserved that limitation rather than treating the run as a complete reproduction.",
"Passing extension checks and reproducing every paper result are separate claims."]),
('Slide 27 Where to Find the Code',[
"The code is organized around four main files.",
"The physical response model is in grid_model.py. This is where the optional service weights and critical-first objective are implemented.",
"The reusable extension logic is in extensions.py. It contains the protection functions, outage tables and timing calculations.",
"The experiment definitions are in run_extensions.py. This file controls scenario splits, parameter sweeps and result summaries.",
"Finally, audit_extensions.py contains the independent checks and the numerical-precision comparison. All extension responses use the group repository's SciPy defender."]),
('Slide 28 Demonstration',[
"The offline demonstration provides three ways to inspect the results.",
"For protection, we can see which lines are protected and which replacement attack becomes worst afterward.",
"For timing, we can compare the same selected attack under even, early and late response windows.",
"For priority service, we can inspect both a selected positive example and a selected adverse example.",
"The demonstration displays saved traces. It does not run optimization in the browser, which makes it convenient to use during the presentation. The detailed measurements remain available in the CSV and JSON files."]),
('Slide 29 Conclusions',[
"Our conclusions are mixed, and that is important.",
"The protection extension finds useful plans, especially on case 14, but does not consistently beat simpler baselines under unseen conditions.",
"The timing extension shows that failure spacing changes the final loss and the snapshot energy proxy. These metrics should be interpreted separately.",
"The priority extension can preserve selected critical demand, but may increase other losses or make a decision that harms later critical service.",
"Therefore, our contribution is not a claim that every extension always improves resilience. It is an implemented and evaluated investigation of when these changes help, when they tie, and when they fail."]),
('Slide 30 Future Work',[
"The next steps follow directly from these findings.",
"Protection should be designed across a wider range of demand and response-speed conditions. Critical-service evaluation should search more strongly for attacks targeting critical demand and test more priority-group assignments.",
"We should also examine how generator restart assumptions and dispatch tie-breaking affect the conclusions.",
"Finally, a more realistic timing model would include continuous ramps, load restoration and storage, followed by validation using more detailed electrical models.",
"This completes our second-part contribution: critique, three implemented extensions, evaluation, demonstrations, conclusions and future work. Thank you."])
])]


def build():
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=Inches(.65);section.bottom_margin=Inches(.65)
    section.left_margin=Inches(.75);section.right_margin=Inches(.75)
    for name,size,bold in [('Normal',11.5,False),('Title',21,True),('Heading 1',16,True),('Heading 2',12,True)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.line_spacing=1.08
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'Arial')
        if name.startswith('Heading'):style.paragraph_format.keep_with_next=True
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
    doc.add_paragraph('SEQUIN Presentation Script',style='Title')
    doc.add_paragraph('Fifteen minutes divided between three speakers, with approximately five minutes each. This script follows the new slides 15 to 30 in the combined presentation.')
    p=doc.add_paragraph();p.add_run('Easiest role  ').bold=True;p.add_run('Part Three is the least technical. It covers verification, code locations, demonstrations, conclusions and future work.')
    p=doc.add_paragraph('Speak slowly, point to the figures, and pause briefly between slides. Rehearse once with a timer and adjust pauses to keep each part near five minutes.')
    p.paragraph_format.space_after=Pt(10)
    split_at=[2,2,3]
    for part_index,(title,meta,slides) in enumerate(PARTS):
        if part_index:doc.add_page_break()
        doc.add_paragraph(title,style='Heading 1')
        p=doc.add_paragraph(meta);p.paragraph_format.space_after=Pt(9)
        for j,(heading,paragraphs) in enumerate(slides):
            if j==split_at[part_index]:
                doc.add_page_break()
                doc.add_paragraph(title+' Continued',style='Heading 1')
                doc.add_paragraph(meta)
            doc.add_paragraph(heading,style='Heading 2')
            for text in paragraphs:doc.add_paragraph(text)
        if part_index==2:
            p=doc.add_paragraph('Move to the original Thank You slide.');p.runs[0].italic=True
    foot=section.footer.paragraphs[0];foot.alignment=2
    run=foot.add_run();run.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
    doc.core_properties.title='SEQUIN Presentation Script'
    doc.core_properties.subject='Fifteen minute script for three speakers covering slides 15 to 30'
    doc.core_properties.author=''
    # Word's default Title style can carry a themed blue bottom rule.
    # This speaking script uses typography and whitespace only.
    for element in (doc.element,doc.styles.element):
        for border in list(element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    print(OUT)
    print('Spoken words by part:',[sum(len(re.findall(r"\b[\w]+\b",p)) for _,ps in slides for p in ps) for _,_,slides in PARTS])


if __name__=='__main__':build()
